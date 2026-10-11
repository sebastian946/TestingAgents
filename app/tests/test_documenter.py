"""Acceptance tests for WTA-19: the Excel report.

The workbook is built from plain data and re-opened with openpyxl; the worker test runs a
whole job against a SQLite file (fake crawl and Designer from conftest.py) and checks that
the Summary totals match the database. Opening in real Excel is checked separately (README).
"""
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import sessionmaker

from agents.documenter import ReportJob, ReportPage, ReportScenario, build_workbook, report_file
from agents.explorer import CrawledPage
from db import db_crud
from db.conn import Base
from db.models_db.models_db import Job, JobStatus, Scenario
from agents import graph
from worker import tasks

UTC = timezone.utc
JOB = ReportJob(uuid.uuid4(), "https://shop.example/", "Online store. Checkout is critical.", "running",
                datetime(2026, 10, 11, 14, 0, tzinfo=UTC), None)
PAGES = [
    ReportPage(1, "https://shop.example/login", "Sign in", "login"),
    ReportPage(2, "https://shop.example/checkout", "Checkout", "checkout"),
    ReportPage(3, "https://shop.example/about", "About us", "content"),  # no scenarios
]
SCENARIOS = [
    ReportScenario("SC-001", 1, "Login works with valid credentials", ["Open the page", "Type <valid user>", "Click Login"],
                   "The home page opens", "critical", "functional"),
    ReportScenario("SC-002", 1, "Login fails with a wrong password", ["Open the page", "Type a wrong password"],
                   "An error is shown", "high", "negative"),
    ReportScenario("SC-003", 2, "Order is placed with a valid card", ["Fill the form", "Click Place order"],
                   "A confirmation appears", "critical", "functional"),
]


def _save(tmp_path: Path, scenarios=SCENARIOS, pages=PAGES, job=JOB):
    path = tmp_path / "report.xlsx"
    build_workbook(job, pages, scenarios, generated_at=datetime(2026, 10, 11, 15, 30, tzinfo=UTC)).save(path)
    return load_workbook(path)


def _summary_values(ws) -> dict:
    """Label in column A -> value in column B (facts) or D (breakdown counts)."""
    out = {}
    for row in ws.iter_rows(min_row=1, max_col=4):
        label = row[0].value
        if isinstance(label, str):
            out.setdefault(label, row[1].value if row[3].value is None else row[3].value)
    return out


def test_two_sheets_summary_first(tmp_path):
    wb = _save(tmp_path)
    assert wb.sheetnames == ["Summary", "Scenarios"]
    assert wb.active.title == "Summary"


def test_scenarios_sheet_has_the_ticket_columns_and_one_step_per_line(tmp_path):
    ws = _save(tmp_path)["Scenarios"]
    headers = [c.value for c in ws[1]]
    assert headers[:8] == ["ID", "Page", "Page type", "Title", "Steps", "Expected result", "Priority", "Category"]
    assert headers[8:] == ["Result", "Notes"]  # left empty for whoever runs the tests
    first = [c.value for c in ws[2]]
    assert first[0] == "SC-001" and first[1] == "https://shop.example/login"
    assert first[4] == "1. Open the page\n2. Type <valid user>\n3. Click Login"
    assert ws["E2"].alignment.wrap_text
    assert ws.max_row == 1 + len(SCENARIOS)


def test_basic_formatting_bold_colored_headers_widths_filters_frozen_header(tmp_path):
    ws = _save(tmp_path)["Scenarios"]
    header = ws["A1"]
    assert header.font.bold and header.fill.fgColor.rgb.endswith("1F4E78")
    assert ws.auto_filter.ref == "A1:J4"
    assert ws.freeze_panes == "B2"
    assert ws.column_dimensions["E"].width >= 50  # steps get room
    assert ws.row_dimensions[2].height >= 45  # 3 steps -> at least 3 lines tall
    assert ws["G2"].fill.fgColor.rgb.endswith("FFC7CE")  # critical is highlighted
    assert any("Pass" in dv.formula1 for dv in ws.data_validations.dataValidation)


def test_summary_totals_match_the_scenarios(tmp_path):
    values = _summary_values(_save(tmp_path)["Summary"])
    assert values["Site"] == JOB.url
    assert values["Pages explored"] == 3
    assert values["Test scenarios"] == 3
    assert values["Pages without scenarios"] == 1
    assert (values["critical"], values["high"], values["medium"], values["low"]) == (2, 1, 0, 0)  # every level shown
    assert (values["functional"], values["negative"], values["security"]) == (2, 1, 0)


def test_summary_explains_itself(tmp_path):
    ws = _save(tmp_path)["Summary"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "How to use this report" in text
    assert "A core flow is blocked" in text  # what "critical" means
    assert "<angle brackets> are placeholders" in text


def test_untrusted_text_cannot_become_a_formula_or_break_the_file(tmp_path):
    evil_page = ReportPage(1, "https://shop.example/x", '=HYPERLINK("http://evil.example","Click")', "content")
    evil = ReportScenario("SC-001", 1, "=1+1\x07 with a bell character", ["=cmd|'/C calc'!A0"], "ok", "low", "usability")
    wb = _save(tmp_path, scenarios=[evil], pages=[evil_page])
    ws = wb["Scenarios"]
    assert ws["D2"].value == "=1+1 with a bell character" and ws["D2"].data_type == "s"
    assert ws["E2"].value == "1. =cmd|'/C calc'!A0"
    summary_title = next(c for row in wb["Summary"].iter_rows() for c in row if c.value and str(c.value).startswith("=HYPER"))
    assert summary_title.data_type == "s"


def test_timezone_aware_dates_are_written_as_utc(tmp_path):
    job = ReportJob(uuid.uuid4(), "https://a.example/", None, "done", None, None)
    path = tmp_path / "r.xlsx"
    build_workbook(job, [], [], generated_at=datetime(2026, 10, 11, 10, 0, tzinfo=timezone(timedelta(hours=-5)))).save(path)
    values = _summary_values(load_workbook(path)["Summary"])
    assert values["Report generated (UTC)"] == datetime(2026, 10, 11, 15, 0)
    assert values["App description"] == "(not provided)"


def test_empty_job_still_produces_a_valid_report(tmp_path):
    ws = _save(tmp_path, scenarios=[], pages=[])["Scenarios"]
    assert ws.max_row == 1 and ws.auto_filter.ref == "A1:J2"


# ---------------------------------------------------------------- end to end, worker

@pytest.fixture
def sessions(tmp_path: Path, monkeypatch, fake_designer):
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False)
    monkeypatch.setattr(tasks, "local_session", factory)
    monkeypatch.setattr(graph, "local_session", factory)
    monkeypatch.setattr(graph.settings, "reports_dir", str(tmp_path / "reports"))

    def crawl(url, on_page=None, **_kwargs):
        pages = [CrawledPage(url=f"{url}p{i}", title=f"Page {i}", status_code=200, depth=min(i, 1)) for i in range(3)]
        for page in pages:
            on_page(page)
        return pages

    monkeypatch.setattr(graph, "crawl", crawl)
    yield factory
    engine.dispose()


def test_a_finished_job_has_its_report_and_the_totals_match_the_database(sessions, fake_designer, tmp_path):
    fake_designer.per_page = 3
    fake_designer.fail_urls = {"https://shop.example/p2": "declined"}
    with sessions() as db:
        job_id = db_crud.create_job(db, url="https://shop.example/", description="A shop").id

    tasks.process_url_task(str(job_id))

    with sessions() as db:
        job = db.get(Job, job_id)
        db_priorities = dict(db.execute(
            select(Scenario.priority, func.count()).where(Scenario.job_id == job_id).group_by(Scenario.priority)
        ).all())
        db_total = db.scalar(select(func.count()).select_from(Scenario).where(Scenario.job_id == job_id))
    path = report_file(tmp_path / "reports", job_id)
    assert job.status is JobStatus.DONE and job.report_path == str(path) and path.exists()

    wb = load_workbook(path)
    values = _summary_values(wb["Summary"])
    assert values["Test scenarios"] == db_total == 6 == wb["Scenarios"].max_row - 1
    assert values["Pages explored"] == 3 and values["Pages without scenarios"] == 1
    for priority in ("critical", "high", "medium", "low"):
        assert values[priority] == db_priorities.get(priority, 0)
    assert not list(path.parent.glob("*.tmp"))  # atomic save left nothing behind


def test_reprocessing_replaces_the_previous_report(sessions, fake_designer, tmp_path):
    with sessions() as db:
        job_id = db_crud.create_job(db, url="https://shop.example/").id
    tasks.process_url_task(str(job_id))
    path = report_file(tmp_path / "reports", job_id)
    first = path.stat().st_mtime_ns

    fake_designer.per_page = 1
    tasks.process_url_task(str(job_id))
    assert path.stat().st_mtime_ns != first
    assert load_workbook(path)["Scenarios"].max_row - 1 == 3
