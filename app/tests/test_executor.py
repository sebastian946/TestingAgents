"""The Executor (WTA-33): browser tools and guardrails in Chromium against a local site, the
tool-use loop with a scripted Anthropic client, the run task in fake mode, and the API.
"""
import threading
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from agents import executor
from agents.executor import BrowserTools, RefusedAction, ScenarioToRun, execute_scenario
from agents.explorer import PlaywrightFetcher
from config.variables import settings
from db import db_crud
from db.conn import Base, get_db
from db.models_db.models_db import TestRun as RunModel
from main import app
from routes import routes, run_routes
from tests.conftest import sign_in
from worker import run_tests, tasks

from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer


def in_thread(fn, *args):
    """Run a worker task outside the test thread: the TestClient leaves an asyncio loop
    there and Playwright's sync API refuses to start inside one (as RQ runs it, too)."""
    box = {}

    def target():
        try:
            box["value"] = fn(*args)
        except BaseException as exc:  # noqa: BLE001
            box["error"] = exc

    t = threading.Thread(target=target)
    t.start()
    t.join()
    if "error" in box:
        raise box["error"]
    return box.get("value")

PAGES = {
    "/": "<title>Shop</title><h1>Shop</h1><a href='/contact'>Contact us</a><a href='/account/delete'>Delete account</a><a href='/checkout'>Checkout</a>",
    "/contact": """<title>Contact</title><form onsubmit="event.preventDefault();document.getElementById('msg').textContent=
        document.querySelector('[name=email]').value.includes('@') ? 'Thanks, we got it' : 'Email is invalid'">
        <label for=e>Email</label><input id=e name=email><label>Message <textarea name=message></textarea></label>
        <button type=submit>Send message</button></form><p id=msg role=status></p>""",
    "/checkout": "<title>Checkout</title><button>Pay now</button>",
    "/account/delete": "<title>Deleted</title>",
}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *_):
        pass

    def do_GET(self):
        body = PAGES.get(self.path.split("?")[0])
        data = f"<html><body>{body or 'nope'}</body></html>".encode()
        self.send_response(200 if body else 404)
        self.send_header("Content-Type", "text/html")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


@pytest.fixture(scope="module")
def site():
    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_port}"
    server.shutdown()


@pytest.fixture(scope="module")
def browser():
    with PlaywrightFetcher() as fetcher:
        yield fetcher


@pytest.fixture
def tools(browser, site):
    page = browser.new_page()
    yield BrowserTools(page, "127.0.0.1")
    page.close()


# ----------------------------------------------------------------- tools and guardrails

def test_tools_navigate_read_fill_click(tools, site):
    assert tools.navigate(f"{site}/contact").startswith("Opened")
    inventory = executor.json.loads(tools.read_page())
    assert inventory["title"] == "Contact" and {i["name"] for i in inventory["inputs"]} == {"email", "message"}
    assert [b["text"] for b in inventory["buttons"]] == ["Send message"]
    tools.fill("Email", "not-an-email")
    tools.fill("Message", "hello")
    tools.click("Send message")
    assert "Email is invalid" in executor.json.loads(tools.read_page())["text"]
    tools.fill("email", "qa@example.com")  # by name this time
    tools.press("Enter")
    assert "Thanks, we got it" in executor.json.loads(tools.read_page())["text"]


def test_tools_refuse_other_hosts_and_destructive_actions(tools, site):
    with pytest.raises(RefusedAction, match="not on the site"):
        tools.navigate("https://evil.example/")
    with pytest.raises(RefusedAction, match="destructive or paid"):
        tools.navigate(f"{site}/account/delete")
    tools.navigate(f"{site}/")
    with pytest.raises(RefusedAction, match="destructive or paid"):
        tools.click("Delete account")
    with pytest.raises(RefusedAction, match="destructive or paid"):
        tools.click("Checkout")
    with pytest.raises(RefusedAction, match="Element not found"):
        tools.click("Nonexistent button")
    with pytest.raises(RefusedAction, match="unknown key"):
        tools.press("rm -rf")


# ----------------------------------------------------------------- the loop

def _block(type_, **kw):
    return SimpleNamespace(type=type_, **kw)


def _tool_use(name, **input_):
    return _block("tool_use", id=f"tu_{name}_{uuid.uuid4().hex[:4]}", name=name, input=input_)


def _response(*content):
    return SimpleNamespace(content=list(content), model="claude-sonnet-5-5",
                           usage=SimpleNamespace(input_tokens=500, output_tokens=50, cache_read_input_tokens=0, cache_creation_input_tokens=0),
                           stop_reason="tool_use")


class ScriptedClient:
    def __init__(self, *script):
        self.script, self.calls = list(script), []
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kwargs):
        self.calls.append({**kwargs, "messages": list(kwargs["messages"])})  # a snapshot: the loop mutates the list
        return self.script.pop(0) if len(self.script) > 1 else self.script[0]


SCENARIO = ScenarioToRun(1, "SC-001", "Contact form rejects an invalid email", ["Open the contact page", "Type 'bad' in Email", "Click Send message"],
                         "An error says the email is invalid", page_url="")


def test_loop_executes_tools_and_records_the_verdict(tools, site, monkeypatch, tmp_path):
    monkeypatch.setattr(settings, "designer_mode", "anthropic")
    scenario = ScenarioToRun(**{**SCENARIO.__dict__, "page_url": f"{site}/contact"})
    client = ScriptedClient(
        _response(_tool_use("read_page")),
        _response(_tool_use("fill", target="Email", text="bad"), _tool_use("click", target="Send message")),
        _response(_tool_use("read_page")),
        _response(_block("text", text="Done."), _tool_use("finish", outcome="pass", reason="The page shows 'Email is invalid'.")),
    )
    calls = []
    result = execute_scenario(scenario, tools, client=client, model="claude-sonnet-5-5", on_call=calls.append,
                              screenshot_path=tmp_path / "SC-001.png")

    assert result.outcome == "pass" and "invalid" in result.reason and result.calls == 4
    assert [s.tool for s in result.steps] == ["read_page", "fill", "click", "read_page", "finish"]
    # the full page went back to the model; the step log keeps a prefix for the record
    assert "Email is invalid" in client.calls[3]["messages"][-1]["content"][0]["content"]
    assert Path(result.screenshot_path).is_file()
    assert len(calls) == 4 and all(c.model == "claude-sonnet-5-5" for c in calls)
    first = client.calls[0]
    assert first["model"] == "claude-sonnet-5-5" and first["output_config"] == {"effort": settings.executor_effort}
    assert [t["name"] for t in first["tools"]] == ["navigate", "read_page", "click", "fill", "press", "finish"]
    assert "SC-001" in first["messages"][0]["content"] and "<expected_result>" in first["messages"][0]["content"]
    # tool results went back as tool_result blocks for the right ids
    second = client.calls[1]["messages"]
    assert second[-1]["role"] == "user" and second[-1]["content"][0]["type"] == "tool_result"


def test_loop_refusals_reach_the_model_and_the_step_cap_blocks(tools, site, monkeypatch):
    monkeypatch.setattr(settings, "designer_mode", "anthropic")
    scenario = ScenarioToRun(**{**SCENARIO.__dict__, "page_url": f"{site}/"})
    client = ScriptedClient(_response(_tool_use("click", target="Delete account")))  # forever
    result = execute_scenario(scenario, tools, client=client, model="claude-sonnet-5-5", max_steps=3)

    assert result.outcome == "blocked" and "Step limit of 3" in result.reason
    assert all("Refused" in s.result for s in result.steps) and len(result.steps) == 3
    assert "Refused" in client.calls[1]["messages"][-1]["content"][0]["content"]


def test_loop_without_tool_calls_is_blocked_and_bad_outcome_is_normalized(tools, site, monkeypatch):
    monkeypatch.setattr(settings, "designer_mode", "anthropic")
    scenario = ScenarioToRun(**{**SCENARIO.__dict__, "page_url": f"{site}/"})
    prose = ScriptedClient(_response(_block("text", text="I think it passes.")))
    assert execute_scenario(scenario, tools, client=prose, model="claude-sonnet-5-5").outcome == "blocked"
    weird = ScriptedClient(_response(_tool_use("finish", outcome="maybe", reason="?")))
    assert execute_scenario(scenario, tools, client=weird, model="claude-sonnet-5-5").outcome == "blocked"


def test_a_page_that_does_not_open_is_blocked_without_calling_the_model(tools, monkeypatch):
    monkeypatch.setattr(settings, "designer_mode", "anthropic")
    scenario = ScenarioToRun(**{**SCENARIO.__dict__, "page_url": "https://evil.example/x"})
    client = ScriptedClient(_response())
    result = execute_scenario(scenario, tools, client=client)
    assert result.outcome == "blocked" and "not on the site" in result.reason and client.calls == []


# ----------------------------------------------------------------- the run task + API

@pytest.fixture
def env(tmp_path: Path, monkeypatch, fake_designer, site):
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False)
    from agents import graph
    for module in (tasks, graph, run_tests):
        monkeypatch.setattr(module, "local_session", factory)
    monkeypatch.setattr(graph.settings, "reports_dir", str(tmp_path / "reports"))
    monkeypatch.setattr(settings, "default_plan", "pro")
    monkeypatch.setattr(settings, "designer_mode", "fake")  # the Executor passes any page that loads
    monkeypatch.setattr(settings, "credentials_key", "u0cZQ6bQ6o9aRtd8cP9yq1lJ5jx0X2uY1hL1gmB2v0E=")
    monkeypatch.setattr(routes, "validate_public_url", lambda url: url)
    monkeypatch.setattr(run_routes, "validate_public_url", lambda url: url)
    monkeypatch.setattr(routes, "add_new_job_to_queue", lambda job_id: None)
    enqueued = []
    monkeypatch.setattr(run_routes, "enqueue_test_run", enqueued.append)
    from agents.explorer import CrawledPage

    def crawl(url, on_page=None, **_kwargs):
        pages = [CrawledPage(url=f"{site}/", title="Shop", status_code=200, depth=0, page_type="content"),
                 CrawledPage(url=f"{site}/contact", title="Contact", status_code=200, depth=1, page_type="form")]
        for page in pages:
            on_page(page)
        return pages

    monkeypatch.setattr(graph, "crawl", crawl)

    def override_db():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    client = TestClient(app)
    client.user_id = sign_in(client)
    yield {"factory": factory, "client": client, "enqueued": enqueued, "site": site}
    app.dependency_overrides.clear()
    engine.dispose()


def _done_job(env) -> str:
    job_id = env["client"].post("/jobs", json={"url": f"{env['site']}/"}).json()["id"]
    tasks.process_url_task(job_id)
    assert env["client"].get(f"/jobs/{job_id}").json()["status"] == "done"
    return job_id


def test_start_run_needs_a_test_environment_a_done_job_and_the_pro_plan(env, monkeypatch):
    client = env["client"]
    job_id = _done_job(env)
    assert client.post(f"/jobs/{job_id}/runs", json={"test_environment": False}).status_code == 400
    queued = client.post("/jobs", json={"url": f"{env['site']}/"}).json()["id"]
    assert client.post(f"/jobs/{queued}/runs", json={"test_environment": True}).status_code == 409

    monkeypatch.setattr(settings, "default_plan", "starter")
    r = client.post(f"/jobs/{job_id}/runs", json={"test_environment": True})
    assert r.status_code == 402 and "Pro plan" in r.json()["detail"]["message"]
    monkeypatch.setattr(settings, "default_plan", "pro")

    r = client.post(f"/jobs/{job_id}/runs", json={"test_environment": True})
    assert r.status_code == 202 and r.json()["status"] == "queued" and env["enqueued"] == [r.json()["id"]]
    assert client.post(f"/jobs/{job_id}/runs", json={"test_environment": True}).status_code == 409  # one at a time
    assert client.get(f"/jobs/{job_id}/runs").json()[0]["id"] == r.json()["id"]
    assert client.get("/billing/me").json()["runs_this_month"] == 1


def test_the_monthly_run_allowance_is_enforced(env, monkeypatch):
    client = env["client"]
    job_id = _done_job(env)
    from config import plans
    monkeypatch.setitem(plans.PLANS, "pro", plans.Plan("pro", "Pro", 40, 49, test_runs_per_month=1))
    assert client.post(f"/jobs/{job_id}/runs", json={"test_environment": True}).status_code == 202
    with env["factory"]() as db:
        db_crud.finish_run(db, 1)
    r = client.post(f"/jobs/{job_id}/runs", json={"test_environment": True})
    assert r.status_code == 402 and r.json()["detail"] == {"message": r.json()["detail"]["message"], "plan": "pro", "used": 1, "limit": 1}
    assert client.get("/billing/me").json()["runs_remaining"] == 0


def test_the_run_task_executes_persists_and_adds_the_results_sheet(env):
    """DESIGNER_MODE=fake (conftest): the Executor passes pages that load, no model."""
    client = env["client"]
    job_id = _done_job(env)
    run_id = client.post(f"/jobs/{job_id}/runs", json={
        "test_environment": True,
        "login": {"login_url": f"{env['site']}/contact", "username": "qa", "password": "x", "authorized": True},
    }).json()["id"]
    with env["factory"]() as db:
        assert db.get(RunModel, run_id).login_credentials  # stored encrypted until the run ends

    in_thread(tasks.run_tests_task, run_id)

    run = client.get(f"/jobs/{job_id}/runs/{run_id}").json()
    assert run["status"] == "failed" and "No password field" in run["error"]  # /contact has no password field
    with env["factory"]() as db:
        assert db.get(RunModel, run_id).login_credentials is None

    run_id = client.post(f"/jobs/{job_id}/runs", json={"test_environment": True}).json()["id"]
    in_thread(tasks.run_tests_task, run_id)
    run = client.get(f"/jobs/{job_id}/runs/{run_id}").json()
    assert run["status"] == "done" and run["total"] == run["passed"] > 0 and run["failed"] == 0
    assert all(r["outcome"] == "pass" and r["has_screenshot"] for r in run["results"])
    shot = client.get(f"/jobs/{job_id}/runs/{run_id}/results/{run['results'][0]['id']}/screenshot")
    assert shot.status_code == 200 and shot.headers["content-type"] == "image/png"

    report = client.get(f"/jobs/{job_id}/report")
    wb = load_workbook(__import__("io").BytesIO(report.content))
    assert wb.sheetnames == ["Summary", "Scenarios", "Results"]
    rows = list(wb["Results"].iter_rows(min_row=6, values_only=True))
    assert len(rows) == run["total"] and all(r[2] == "Pass" for r in rows)

    other = TestClient(app)
    sign_in(other, email="other@example.com")
    assert other.get(f"/jobs/{job_id}/runs/{run_id}").status_code == 404


def test_run_failures_are_recorded_not_lost(env, monkeypatch):
    client = env["client"]
    job_id = _done_job(env)
    run_id = client.post(f"/jobs/{job_id}/runs", json={"test_environment": True}).json()["id"]

    def boom(*a, **k):
        raise RuntimeError("browser crashed")

    monkeypatch.setattr(run_tests, "execute_scenario", boom)
    with pytest.raises(RuntimeError):
        in_thread(tasks.run_tests_task, run_id)
    run = client.get(f"/jobs/{job_id}/runs/{run_id}").json()
    assert run["status"] == "failed" and "browser crashed" in run["error"]
