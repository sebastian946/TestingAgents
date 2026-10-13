"""End-to-end integration test: API -> Redis queue -> worker -> graph -> Excel (WTA-24).

Run it with one command, from app/ (Docker Desktop must be running):

    uv run pytest -m integration -v

What is real and what is not:
- Real: PostgreSQL and Redis from docker-compose, the FastAPI app and the RQ worker as
  separate processes, the LangGraph pipeline, the Playwright/Chromium crawl, the page
  classifier, the persistence and the Excel report.
- Local: the site under test is served from 127.0.0.1 by this file, so the test needs no
  internet. The anti-SSRF check lets 127.0.0.1 through only because SSRF_ALLOWED_HOSTS says so.
- Fake: the LLM. DESIGNER_MODE=fake makes the Designer build scenarios from the page
  inventory with fixed rules. A configuration switch rather than a mock because the worker
  is another process: a monkeypatch here would not reach it. No API key, no cost, and the
  same scenarios on every run.

Isolation: its own database (testing_agents_it), its own RQ queue name per run (a worker you
left running on the default queue cannot steal its jobs) and a temporary reports folder.

Every wait is a poll with a timeout (`wait_until`), never a fixed sleep: services, the API
and the jobs all finish when they finish, and the test moves on as soon as they do.
"""
import io
import json
import os
import socket
import subprocess
import sys
import threading
import time
import urllib.error
import urllib.request
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, Callable

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.integration

APP_DIR = Path(__file__).resolve().parents[1]
IT_DB_NAME = "testing_agents_it"
SERVICE_TIMEOUT = 90  # docker services and the API coming up
JOB_TIMEOUT = 180  # a full job: Chromium crawl of a few local pages + fake design + report


# ------------------------------------------------------------------------- polling

def wait_until(check: Callable[[], Any], timeout: float, what: str, interval: float = 0.5) -> Any:
    """Call `check` until it returns something truthy, or fail after `timeout` seconds.

    Exceptions inside `check` count as "not yet" (a service still starting refuses
    connections); the last one is reported if time runs out.
    """
    deadline = time.monotonic() + timeout
    last: Any = None
    while time.monotonic() < deadline:
        try:
            result = check()
            if result:
                return result
            last = result
        except Exception as exc:  # noqa: BLE001 - "not ready yet" can look like anything
            last = exc
        time.sleep(interval)
    pytest.fail(f"Timed out after {timeout:.0f}s waiting for {what}. Last result: {last!r}")


# --------------------------------------------------------------------- HTTP helpers

def http(method: str, url: str, body: dict | None = None) -> tuple[int, dict[str, str], bytes]:
    data = json.dumps(body).encode() if body is not None else None
    request = urllib.request.Request(url, data=data, method=method,
                                     headers={"Content-Type": "application/json"} if data else {})
    try:
        with urllib.request.urlopen(request, timeout=30) as response:
            return response.status, dict(response.headers), response.read()
    except urllib.error.HTTPError as err:
        return err.code, dict(err.headers), err.read()


def get_json(url: str) -> Any:
    status, _, body = http("GET", url)
    assert status == 200, f"GET {url} -> {status}: {body[:300]!r}"
    return json.loads(body)


def free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def kill_tree(process: subprocess.Popen) -> None:
    """Stop a process and its children (on Windows the venv launcher spawns a child)."""
    if process.poll() is not None:
        return
    if os.name == "nt":
        subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], capture_output=True)
    else:
        process.terminate()
    try:
        process.wait(timeout=15)
    except subprocess.TimeoutExpired:
        process.kill()


# ----------------------------------------------------------------- site under test

SITE = {
    "/": """<html><head><title>Acme Store</title></head><body>
        <nav><a href="/login">Sign in</a> <a href="/contact">Contact</a> <a href="/about">About</a></nav>
        <h1>Welcome to Acme</h1></body></html>""",
    "/login": """<html><head><title>Sign in</title></head><body><h1>Sign in</h1>
        <form action="/session" method="post">
          <label for="u">Username</label><input id="u" name="username" required>
          <label for="p">Password</label><input id="p" name="password" type="password" required>
          <button type="submit">Log in</button>
        </form></body></html>""",
    "/contact": """<html><head><title>Contact</title></head><body><h1>Contact us</h1>
        <form action="/send" method="post">
          <label>Your name <input name="name" required></label>
          <label>Email <input name="email" type="email" required></label>
          <label>Message <textarea name="message"></textarea></label>
          <button type="submit">Send message</button>
        </form></body></html>""",
    "/about": """<html><head><title>About</title></head><body><h1>About Acme</h1>
        <p>We sell things.</p><a href="/">Home</a></body></html>""",
}


class SiteHandler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def do_GET(self):
        path = self.path.split("?")[0].rstrip("/") or "/"
        if path == "/robots.txt":
            body, ctype, status = "User-agent: *\nAllow: /\n", "text/plain", 200
        elif path in SITE:
            body, ctype, status = SITE[path], "text/html; charset=utf-8", 200
        else:
            body, ctype, status = "<h1>Not found</h1>", "text/html", 404
        data = body.encode()
        self.send_response(status)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)


@pytest.fixture(scope="module")
def site_url():
    server = ThreadingHTTPServer(("127.0.0.1", 0), SiteHandler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    yield f"http://127.0.0.1:{server.server_port}"
    server.shutdown()


# ------------------------------------------------------------------ the services

@pytest.fixture(scope="module")
def services():
    """PostgreSQL and Redis from docker-compose, started if needed and polled until ready."""
    from config.variables import settings

    compose = ["docker", "compose", "--env-file", "../.env", "-f", "docker-compose.yml"]
    try:
        up = subprocess.run([*compose, "up", "-d", "postgres", "redis"], cwd=APP_DIR,
                            capture_output=True, text=True, timeout=180)
    except (FileNotFoundError, subprocess.TimeoutExpired) as exc:
        pytest.skip(f"Docker is not available: {exc}")
    if up.returncode != 0:
        pytest.skip(f"Could not start docker-compose services (is Docker Desktop running?): {up.stderr.strip()[-300:]}")

    import psycopg
    from redis import Redis

    admin_dsn = (f"host={settings.endpoint} port={settings.db_port} user={settings.db_user} "
                 f"password={settings.db_password} dbname=postgres")
    wait_until(lambda: psycopg.connect(admin_dsn, connect_timeout=2).close() is None, SERVICE_TIMEOUT, "PostgreSQL")
    wait_until(lambda: Redis(host=settings.redis_host, port=settings.redis_port, socket_timeout=2).ping(),
               SERVICE_TIMEOUT, "Redis")

    # A database of its own: the test never touches your development data
    with psycopg.connect(admin_dsn, autocommit=True) as conn:
        exists = conn.execute("SELECT 1 FROM pg_database WHERE datname = %s", (IT_DB_NAME,)).fetchone()
        if not exists:
            conn.execute(f'CREATE DATABASE "{IT_DB_NAME}"')
    return settings


@pytest.fixture(scope="module")
def stack(services, tmp_path_factory):
    """The API and the worker as real processes, configured for the test."""
    settings = services
    run_dir = tmp_path_factory.mktemp("integration")
    api_port = free_port()
    env = {
        **os.environ,
        "DB_NAME": IT_DB_NAME,
        "RQ_QUEUE_NAME": f"integration-{uuid.uuid4().hex[:8]}",  # nobody else listens here
        "REPORTS_DIR": str(run_dir / "reports"),
        "DESIGNER_MODE": "fake",
        "SSRF_ALLOWED_HOSTS": json.dumps(["127.0.0.1"]),
        "LANGSMITH_TRACING": "false",
        "ANTHROPIC_API_KEY": "sk-ant-test-not-a-real-key",
        "ANTHROPIC_BASE_URL": "http://127.0.0.1:9",
        "PYTHONIOENCODING": "utf-8",
    }

    init = subprocess.run([sys.executable, "-m", "db.init_db"], cwd=APP_DIR, env=env, capture_output=True, text=True)
    assert init.returncode == 0, init.stderr[-2000:]
    import psycopg

    dsn = (f"host={settings.endpoint} port={settings.db_port} user={settings.db_user} "
           f"password={settings.db_password} dbname={IT_DB_NAME}")
    with psycopg.connect(dsn, autocommit=True) as conn:  # start every run from empty tables
        conn.execute("TRUNCATE llm_calls, scenarios, pages, job CASCADE")

    logs = {name: open(run_dir / f"{name}.log", "w", encoding="utf-8") for name in ("api", "worker")}
    api = subprocess.Popen([sys.executable, "-m", "uvicorn", "main:app", "--port", str(api_port)],
                           cwd=APP_DIR, env=env, stdout=logs["api"], stderr=subprocess.STDOUT)
    worker = subprocess.Popen([sys.executable, "-m", "worker.worker"],
                              cwd=APP_DIR, env=env, stdout=logs["worker"], stderr=subprocess.STDOUT)
    base = f"http://127.0.0.1:{api_port}"
    try:
        wait_until(lambda: http("GET", f"{base}/Health")[0] == 200, SERVICE_TIMEOUT, "the API /Health")
        wait_until(lambda: "Listening on" in (run_dir / "worker.log").read_text(encoding="utf-8"),
                   SERVICE_TIMEOUT, "the worker to start listening")
        yield {"api": base, "run_dir": run_dir}
    finally:
        for process in (api, worker):
            kill_tree(process)
        for handle in logs.values():
            handle.close()


def _diagnostics(run_dir: Path) -> str:
    parts = []
    for name in ("worker", "api"):
        text = (run_dir / f"{name}.log").read_text(encoding="utf-8", errors="replace")
        parts.append(f"--- {name}.log (last lines) ---\n" + "\n".join(text.splitlines()[-40:]))
    return "\n".join(parts)


def run_job(stack, url: str, language: str = "en", description: str | None = None) -> dict:
    """POST /jobs, then poll GET /jobs/{id} until the job leaves queued/running."""
    status, _, body = http("POST", f"{stack['api']}/jobs", {"url": url, "language": language, "description": description})
    assert status == 202, f"POST /jobs -> {status}: {body!r}"
    job = json.loads(body)
    assert job["status"] == "queued"

    def finished():
        current = get_json(f"{stack['api']}/jobs/{job['id']}")
        return current if current["status"] in ("done", "failed") else None

    try:
        return wait_until(finished, JOB_TIMEOUT, f"job {job['id']} to finish")
    except BaseException:
        print(_diagnostics(stack["run_dir"]))
        raise


# -------------------------------------------------------------------------- tests

def test_a_job_goes_from_the_api_to_a_valid_excel_report(stack, site_url):
    job = run_job(stack, f"{site_url}/", description="A small store")
    assert job["status"] == "done", f"{job}\n{_diagnostics(stack['run_dir'])}"
    api = f"{stack['api']}/jobs/{job['id']}"

    # The crawl followed the links of the local site, and the classifier typed the pages
    pages = get_json(f"{api}/pages")
    by_path = {p["url"].removeprefix(site_url): p for p in pages}
    assert set(by_path) == {"/", "/login", "/contact", "/about"}
    assert job["pages_crawled"] == 4
    assert by_path["/login"]["page_type"] == "login" and by_path["/contact"]["page_type"] == "form"

    # The Designer (fake mode) produced scenarios, and nothing was spent on an LLM
    scenarios = get_json(f"{api}/scenarios")
    assert len(scenarios) == job["total_scenarios"] > 0
    assert {s["scenario_code"] for s in scenarios} == {f"SC-{n:03d}" for n in range(1, len(scenarios) + 1)}
    assert "A wrong password is rejected" in {s["title"] for s in scenarios}
    assert get_json(f"{api}/usage")["calls"] == 0

    # The report downloads as a real .xlsx and matches what the API says
    status, headers, body = http("GET", f"{api}/report")
    assert status == 200
    assert headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'filename="test-report-127.0.0.1-' in headers["content-disposition"]
    wb = load_workbook(io.BytesIO(body))
    assert wb.sheetnames == ["Summary", "Scenarios"]
    rows = list(wb["Scenarios"].iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(scenarios)
    assert {r[0] for r in rows} == {s["scenario_code"] for s in scenarios}
    summary: dict = {}
    for r in wb["Summary"].iter_rows(values_only=True):
        if r[0]:  # first occurrence wins: "Pages explored" is both a fact and a section title
            summary.setdefault(r[0], (r[1], r[3]))
    assert summary["Test scenarios"][0] == len(scenarios)
    assert summary["Pages explored"][0] == 4
    for priority in ("critical", "high", "medium", "low"):
        assert summary[priority][1] == sum(s["priority"] == priority for s in scenarios)


def test_a_spanish_job_gets_a_spanish_report(stack, site_url):
    job = run_job(stack, f"{site_url}/", language="es")
    assert job["status"] == "done" and job["language"] == "es"

    scenarios = get_json(f"{stack['api']}/jobs/{job['id']}/scenarios")
    assert "Una contraseña incorrecta se rechaza" in {s["title"] for s in scenarios}

    status, headers, body = http("GET", f"{stack['api']}/jobs/{job['id']}/report")
    assert status == 200 and 'filename="reporte-pruebas-127.0.0.1-' in headers["content-disposition"]
    assert load_workbook(io.BytesIO(body)).sheetnames == ["Resumen", "Escenarios"]


def test_a_site_with_no_reachable_page_fails_cleanly(stack, site_url):
    job = run_job(stack, f"{site_url}/missing")  # the start page answers 404

    assert job["status"] == "failed"
    assert job["error"] == "No page could be fetched from the given URL."  # the graph's fail node
    status, _, body = http("GET", f"{stack['api']}/jobs/{job['id']}/report")
    assert status == 409 and json.loads(body)["detail"]["status"] == "failed"


def test_an_internal_url_is_still_rejected_by_the_api(stack):
    """The allowlist opens 127.0.0.1 only: the anti-SSRF check still guards everything else."""
    status, _, body = http("POST", f"{stack['api']}/jobs", {"url": "http://169.254.169.254/latest/meta-data/"})
    assert status == 400 and "disallowed" in json.loads(body)["detail"]
