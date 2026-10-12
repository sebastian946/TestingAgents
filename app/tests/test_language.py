"""English / Spanish jobs: the chosen language reaches the Designer and the report.

No network, Redis or LLM: the API runs in-process on SQLite, the Designer is faked or its
client is a stub, and the report is read back with openpyxl.
"""
import io
import uuid
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from agents import designer, graph
from agents.designer import DesignerFatalError, design_scenarios, render_prompt
from agents.designer_eval import LOGIN
from agents.documenter import ReportJob, ReportPage, ReportScenario, build_workbook, download_filename
from agents.explorer import CrawledPage
from db import db_crud
from db.conn import Base, get_db
from main import app
from models.agent_models import DesignerOutput, TestScenario
from routes import routes
from worker import tasks

URL = "https://shop.example/"


@pytest.fixture
def sessions(tmp_path: Path, monkeypatch, fake_designer):
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False)
    for module in (tasks, graph):
        monkeypatch.setattr(module, "local_session", factory)
    monkeypatch.setattr(graph.settings, "reports_dir", str(tmp_path / "reports"))

    def crawl(url, on_page=None, **_kwargs):
        page = CrawledPage(url=f"{url}login", title="Sign in", status_code=200, depth=0, page_type="login")
        on_page(page)
        return [page]

    monkeypatch.setattr(graph, "crawl", crawl)
    yield factory
    engine.dispose()


# ---------------------------------------------------------------------------- API

@pytest.fixture
def client(sessions, monkeypatch):
    monkeypatch.setattr(routes, "validate_public_url", lambda url: url)  # no DNS lookup
    monkeypatch.setattr(routes, "add_new_job_to_queue", lambda job_id: None)  # no Redis

    def override_db():
        with sessions() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    yield TestClient(app)
    app.dependency_overrides.clear()


def test_post_jobs_stores_the_language_and_defaults_to_english(client):
    spanish = client.post("/jobs", json={"url": URL, "language": "es"})
    default = client.post("/jobs", json={"url": URL})
    assert spanish.status_code == 202 and spanish.json()["language"] == "es"
    assert default.status_code == 202 and default.json()["language"] == "en"
    assert client.get(f"/jobs/{spanish.json()['id']}").json()["language"] == "es"


def test_an_unsupported_language_is_rejected(client):
    response = client.post("/jobs", json={"url": URL, "language": "fr"})
    assert response.status_code == 422
    assert "en" in str(response.json()) and "es" in str(response.json())


# ----------------------------------------------------------------------- Designer

def test_the_output_language_goes_in_the_user_message_and_the_system_prompt_is_shared():
    system_en, user_en = render_prompt(LOGIN, None, "v3", "en")
    system_es, user_es = render_prompt(LOGIN, "Una tienda", "v3", "es")
    assert system_en == system_es  # one cached prefix for every language
    assert "<output_language>English</output_language>" in user_en
    assert "<output_language>Spanish" in user_es


def test_older_prompt_versions_ignore_the_language():
    _, user = render_prompt(LOGIN, None, "v2", "es")
    assert "output_language" not in user


def test_an_unsupported_language_is_a_fatal_designer_error():
    with pytest.raises(DesignerFatalError, match="Unsupported language"):
        render_prompt(LOGIN, None, "v3", "fr")


def test_the_request_to_claude_carries_the_language():
    sent = []
    valid = DesignerOutput(scenarios=[TestScenario(
        title="Inicio de sesión correcto", steps=["Abrir la página"], expected_result="Se abre la tienda",
        priority="critical", category="functional",
    )]).model_dump_json()
    usage = SimpleNamespace(input_tokens=1, output_tokens=1, cache_read_input_tokens=0, cache_creation_input_tokens=0)
    response = SimpleNamespace(stop_reason="end_turn", content=[SimpleNamespace(type="text", text=valid)],
                               model="claude-opus-5-5", usage=usage)
    fake = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(create=lambda **kw: sent.append(kw) or response)))

    design_scenarios(LOGIN, None, client=fake, version="v3", language="es")

    assert "<output_language>Spanish" in sent[0]["messages"][0]["content"]


def test_v3_is_the_default_prompt():
    assert designer.settings.designer_prompt_version == "v3"


# --------------------------------------------------------------------------- worker

def test_the_job_language_reaches_the_designer_and_the_report(sessions, fake_designer, tmp_path):
    with sessions() as db:
        job_id = db_crud.create_job(db, url=URL, language="es").id

    tasks.process_url_task(str(job_id))

    assert fake_designer.languages == ["es"]
    wb = load_workbook(tmp_path / "reports" / f"{job_id}.xlsx")
    assert wb.sheetnames == ["Resumen", "Escenarios"]


# --------------------------------------------------------------------------- report

JOB_ES = ReportJob(uuid.uuid4(), URL, None, "done", None, None, "es")
PAGES = [ReportPage(1, f"{URL}login", "Sign in", "login")]
SCENARIOS = [
    ReportScenario("SC-001", 1, "Inicio de sesión correcto", ["Abrir la página"], "Se abre la tienda", "critical", "functional"),
    ReportScenario("SC-002", 1, "Las credenciales no aparecen en la URL", ["Enviar"], "La URL no las muestra", "critical", "security"),
    ReportScenario("SC-003", 1, "Mensaje de error claro", ["Enviar vacío"], "Se ve un error", "medium", "usability"),
]


def _workbook(tmp_path, job=JOB_ES):
    path = tmp_path / "r.xlsx"
    build_workbook(job, PAGES, SCENARIOS, generated_at=datetime(2026, 10, 12, tzinfo=timezone.utc)).save(path)
    return load_workbook(path)


def test_spanish_report_translates_every_label_and_value(tmp_path):
    wb = _workbook(tmp_path)
    summary, sheet = wb["Resumen"], wb["Escenarios"]
    text = " ".join(str(c.value) for row in summary.iter_rows() for c in row if c.value)
    for expected in ("Reporte de escenarios de prueba", "Sitio", "Páginas exploradas", "Escenarios por prioridad",
                     "Qué significa", "Cómo usar este reporte", "(no indicada)", "inicio de sesión"):
        assert expected in text
    assert [c.value for c in sheet[1]] == ["ID", "Página", "Tipo de página", "Título", "Pasos", "Resultado esperado",
                                           "Prioridad", "Categoría", "Resultado", "Notas"]
    assert [c.value for c in sheet[2]][6:8] == ["crítica", "funcional"]
    assert sheet["C2"].value == "inicio de sesión"
    assert sheet["G2"].fill.fgColor.rgb.endswith("FFC7CE")  # colors follow the code, not the word
    assert any("Pasó" in dv.formula1 for dv in sheet.data_validations.dataValidation)


def test_spanish_totals_still_match_the_data(tmp_path):
    summary = _workbook(tmp_path)["Resumen"]
    counts = {row[0].value: row[3].value for row in summary.iter_rows(max_col=4) if row[0].value and row[3].value is not None}
    assert (counts["crítica"], counts["alta"], counts["media"], counts["baja"]) == (2, 0, 1, 0)
    assert (counts["funcional"], counts["seguridad"], counts["usabilidad"]) == (1, 1, 1)


def test_english_report_is_unchanged(tmp_path):
    wb = _workbook(tmp_path, ReportJob(uuid.uuid4(), URL, None, "done", None, None, "en"))
    assert wb.sheetnames == ["Summary", "Scenarios"]
    assert [c.value for c in wb["Scenarios"][2]][6:8] == ["critical", "functional"]


def test_download_name_follows_the_language():
    when = datetime(2026, 10, 12, tzinfo=timezone.utc)
    assert download_filename(URL, when, "es") == "reporte-pruebas-shop.example-2026-10-12.xlsx"
    assert download_filename(URL, when, "en") == "test-report-shop.example-2026-10-12.xlsx"


def test_report_endpoint_serves_the_spanish_name(client, sessions):
    with sessions() as db:
        job_id = db_crud.create_job(db, url=URL, language="es").id
        db_crud.mark_job_running(db, job_id)
        page = db_crud.create_page(db, job_id, url=f"{URL}login", page_type="login")
        db_crud.create_scenarios(db, job_id, page.id, [{"title": "t", "steps": ["s"], "expected_result": "r",
                                                         "priority": "high", "category": "negative"}])
        db_crud.mark_job_done(db, job_id)

    response = client.get(f"/jobs/{job_id}/report")

    assert response.status_code == 200
    assert 'filename="reporte-pruebas-shop.example-' in response.headers["content-disposition"]
    assert load_workbook(io.BytesIO(response.content)).sheetnames == ["Resumen", "Escenarios"]
