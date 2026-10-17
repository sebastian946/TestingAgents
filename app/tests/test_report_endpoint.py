"""Acceptance tests for WTA-20: GET /jobs/{id}/report downloads the Excel of a finished job.

The API runs in-process (TestClient) on a SQLite file; reports go to a temp folder. No
network, no Redis, no LLM.
"""
import io
import uuid
from datetime import datetime, timezone
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from agents.documenter import XLSX_MEDIA_TYPE, download_filename, generate_report, report_file
from db import db_crud
from db.conn import Base, get_db
from db.models_db.models_db import Job, JobStatus
from main import app
from tests.conftest import sign_in


@pytest.fixture
def env(tmp_path: Path, monkeypatch):
    engine = create_engine(f"sqlite:///{tmp_path / 'test.db'}")
    Base.metadata.create_all(engine)
    factory = sessionmaker(bind=engine, autoflush=False)
    reports = tmp_path / "reports"
    from config.variables import settings
    monkeypatch.setattr(settings, "reports_dir", str(reports))

    def override_db():
        with factory() as db:
            yield db

    app.dependency_overrides[get_db] = override_db
    client = TestClient(app)
    client.user_id = sign_in(client)  # the jobs below belong to this account
    yield factory, reports, client
    app.dependency_overrides.clear()
    engine.dispose()


def _job(factory, user_id, url="https://www.Shop-Example.com:8443/home", status="done", scenarios=2, error=None):
    """Create a job for `user_id` in the given state, with one page and some scenarios."""
    with factory() as db:
        job = db_crud.create_job(db, url=url, user_id=user_id)
        db_crud.mark_job_running(db, job.id)
        page = db_crud.create_page(db, job.id, url=url, title="Home", page_type="login")
        db_crud.create_scenarios(db, job.id, page.id, [
            {"title": f"Scenario {i}", "steps": ["Open", "Click"], "expected_result": "Works",
             "priority": "high", "category": "functional"} for i in range(scenarios)
        ])
        if status == "done":
            db_crud.mark_job_done(db, job.id)
        elif status == "failed":
            db_crud.mark_job_failed(db, job.id, error or "boom")
        elif status == "queued":
            db.get(Job, job.id).status = JobStatus.QUEUED
            db.commit()
        return job.id


def test_a_finished_job_downloads_as_a_real_xlsx_with_a_friendly_name(env):
    factory, reports, client = env
    job_id = _job(factory, client.user_id)
    with factory() as db:
        generate_report(db, job_id, reports)

    response = client.get(f"/jobs/{job_id}/report")

    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    disposition = response.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    assert f'filename="test-report-shop-example.com-{today}.xlsx"' in disposition
    wb = load_workbook(io.BytesIO(response.content))  # the bytes are a valid workbook
    assert wb.sheetnames == ["Summary", "Scenarios"] and wb["Scenarios"].max_row == 3


@pytest.mark.parametrize("state", ["queued", "running"])
def test_a_job_in_progress_answers_409_with_progress_and_retry_after(env, state):
    factory, _, client = env
    job_id = _job(factory, client.user_id, status=state)

    response = client.get(f"/jobs/{job_id}/report")

    assert response.status_code == 409
    assert response.headers["retry-after"] == "15"
    detail = response.json()["detail"]
    assert detail["status"] == state
    assert "not ready yet" in detail["message"] and "1 page explored, 2 scenarios so far" in detail["message"]


def test_a_failed_job_answers_409_explaining_why(env):
    factory, _, client = env
    job_id = _job(factory, client.user_id, status="failed", error="The Anthropic API key is missing or invalid.")

    response = client.get(f"/jobs/{job_id}/report")

    assert response.status_code == 409 and "retry-after" not in response.headers
    detail = response.json()["detail"]
    assert detail["status"] == "failed"
    assert "has no report" in detail["message"] and "API key is missing" in detail["message"]


def test_an_unknown_job_is_404(env):
    _, _, client = env
    response = client.get(f"/jobs/{uuid.uuid4()}/report")
    assert response.status_code == 404 and "not found" in response.json()["detail"]


def test_another_users_job_is_404_too(env):
    """Ownership: a job id that exists but belongs to someone else looks like it does not."""
    factory, _, client = env
    job_id = _job(factory, client.user_id)
    other = TestClient(app)
    sign_in(other, email="someone-else@example.com")
    assert other.get(f"/jobs/{job_id}/report").status_code == 404
    assert other.get(f"/jobs/{job_id}").status_code == 404
    assert other.delete(f"/jobs/{job_id}").status_code == 404
    assert client.get(f"/jobs/{job_id}").status_code == 200  # still there for its owner


def test_a_done_job_without_its_file_gets_the_report_rebuilt(env):
    factory, reports, client = env
    job_id = _job(factory, client.user_id, scenarios=3)  # done, but no report was ever written
    assert not report_file(reports, job_id).exists()

    response = client.get(f"/jobs/{job_id}/report")

    assert response.status_code == 200
    assert load_workbook(io.BytesIO(response.content))["Scenarios"].max_row == 4
    with factory() as db:
        assert db_crud.get_job(db, job_id).report_path == str(report_file(reports, job_id))


@pytest.mark.parametrize("url, when, expected", [
    ("https://www.saucedemo.com/", datetime(2026, 10, 11, 23, 30, tzinfo=timezone.utc), "test-report-saucedemo.com-2026-10-11.xlsx"),
    ("https://shop.example.com:8443/a?b=c", datetime(2026, 1, 2), "test-report-shop.example.com-2026-01-02.xlsx"),
    ("https://xn--caf-dma.example/", datetime(2026, 1, 2), "test-report-xn--caf-dma.example-2026-01-02.xlsx"),
    ("not a url", datetime(2026, 1, 2), "test-report-site-2026-01-02.xlsx"),
])
def test_download_filename(url, when, expected):
    assert download_filename(url, when) == expected
